"""Plan files: CSV, xlsx and run log (RULEBOOK Section 12). Rounding happens only here."""
from __future__ import annotations

import csv
import io
from decimal import ROUND_HALF_UP, Decimal
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .engine import COLUMNS, RunResult

DP2 = {"avg4", "avg12", "forward_demand", "cover_weeks", "cover_after"}
DP4 = {"trend"}

FILL = {
    "STOCKOUT_RISK": PatternFill("solid", fgColor="FDE2E1"),
    "OVERSTOCK_RISK": PatternFill("solid", fgColor="FDEFD3"),
    "JUDGMENT_ZONE": PatternFill("solid", fgColor="E8E8ED"),
}
HUMAN_FILL = PatternFill("solid", fgColor="FFF4B3")
ARIAL = Font(name="Arial", size=10)
ARIAL_BOLD = Font(name="Arial", size=10, bold=True)


def fmt(col: str, v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "True" if v else "False"
    if col in DP2:
        return _round_str(v, "0.01")
    if col in DP4:
        return _round_str(v, "0.0001")
    return str(v)


def _round_str(v: float, q: str) -> str:
    # [IMPL] half-up on the shortest decimal repr, so 188.125 prints as 188.13
    return str(Decimal(repr(v)).quantize(Decimal(q), rounding=ROUND_HALF_UP))


def plan_csv(res: RunResult) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(COLUMNS)
    for r in res.rows:
        w.writerow([fmt(c, r[c]) for c in COLUMNS])
    return buf.getvalue().encode("utf-8")


def _xl_value(col: str, v):
    if v is None:
        return None
    if col in DP2:
        return float(_round_str(v, "0.01"))
    if col in DP4:
        return float(_round_str(v, "0.0001"))
    return v


def summary_rows(res: RunResult) -> list[tuple[str, object]]:
    s = res.summary
    return [
        ("Order month / festival", f"{s['order_month']} / {'festival' if s['is_festival'] else 'non-festival'}"),
        ("Inventory as of", s["as_of_date"]),
        ("SKUs: reorder now", s["reorder_now"]),
        ("SKUs: pause next order", s["pause_next_order"]),
        ("SKUs: healthy", s["healthy"]),
        ("SKUs: judgment zone", s["judgment_zone"]),
        ("Total units to order now", s["total_units_now"]),
        ("SKUs needing Raghav review", s["needs_review"]),
        ("Pre-send checks", "PASSED" if res.for_sending else "NOT FOR SENDING"),
    ]


def plan_xlsx(res: RunResult) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = ARIAL_BOLD
    for r in res.rows:
        ws.append([_xl_value(c, r[c]) for c in COLUMNS])
    trend_col = COLUMNS.index("trend") + 1
    for i, r in enumerate(res.rows, start=2):
        fill = FILL.get(r["status"])
        for j, c in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=i, column=j)
            cell.font = ARIAL
            if fill:
                cell.fill = fill
            if j >= 26:  # columns 26-28 are human input
                cell.fill = HUMAN_FILL
        ws.cell(row=i, column=trend_col).number_format = "0.00%"
    for j in range(26, 29):
        ws.cell(row=1, column=j).fill = HUMAN_FILL
    ws.freeze_panes = "A2"
    for j, c in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = max(10, len(c) + 2)

    sm = wb.create_sheet("Summary")
    sm.append(["Metric", "Value"])
    for row in summary_rows(res):
        sm.append(list(row))
    for row in sm.iter_rows():
        for c in row:
            c.font = ARIAL
    for c in sm[1]:
        c.font = ARIAL_BOLD
    sm.column_dimensions["A"].width = 32
    sm.column_dimensions["B"].width = 28
    sm.freeze_panes = "A2"

    ps = wb.create_sheet("Parameters")
    ps.append(["config.yaml"])
    for line in res.config.get("_text", "").splitlines():
        ps.append([line])
    for row in ps.iter_rows():
        for c in row:
            c.font = Font(name="Arial", size=10)
    ps.column_dimensions["A"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_log(res: RunResult, generated_at: datetime | None = None) -> str:
    ts = (generated_at or datetime.now(timezone.utc)).isoformat(timespec="seconds")
    lines = [
        "Freshroots reorder plan run log",
        f"Run timestamp: {ts}",
        f"Run date: {res.run_date.isoformat()}",
        "",
        "Input files:",
    ]
    for name in sorted(res.files):
        f = res.files[name]
        lines.append(f"  {name}: {f.row_count} rows, sha256 {f.sha256}")
    lines += ["", f"Config (sha256 {res.config.get('_sha256', '')}):"]
    lines += [f"  {l}" for l in res.config.get("_text", "").splitlines()]
    lines += ["", "Validation:"]
    for code, passed in res.validation.checks.items():
        lines.append(f"  {code}: {'PASS' if passed else 'FAIL'}")
    for e in res.validation.errors:
        lines.append(f"  {e}")
    if res.fatal:
        lines += ["", "FATAL validation failure: no plan written."]
    else:
        lines += ["", "Pre-send checks:"]
        for code, passed in res.checks.items():
            lines.append(f"  {code}: {'PASS' if passed else 'FAIL'}")
        lines += [f"  {m}" for m in res.check_messages]
        lines += ["", "Summary:"]
        lines += [f"  {k}: {v}" for k, v in summary_rows(res)]
        if not res.for_sending:
            lines.append("PLAN IS NOT FOR SENDING")
    lines += ["", "Warnings:"]
    lines += [f"  {w}" for w in res.warnings] or ["  none"]
    return "\n".join(lines) + "\n"


def write_outputs(res: RunResult, out_dir: str | Path) -> Path:
    """Write output/<order_month>/...; on FATAL only the run log (to out_dir itself)."""
    out = Path(out_dir)
    if res.fatal:
        out.mkdir(parents=True, exist_ok=True)
        (out / f"run_log_{res.run_date.isoformat()}.txt").write_text(run_log(res), encoding="utf-8")
        return out
    month = res.summary["order_month"]
    d = out / month
    d.mkdir(parents=True, exist_ok=True)
    (d / f"reorder_plan_{month}.csv").write_bytes(plan_csv(res))
    (d / f"reorder_plan_{month}.xlsx").write_bytes(plan_xlsx(res))
    (d / f"run_log_{month}.txt").write_text(run_log(res), encoding="utf-8")
    return d
