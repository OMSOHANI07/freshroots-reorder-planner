"""Checks downloaded PO files from the e2e test: only approved quantities, grouped by supplier."""
import csv
import io
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

EXPECTED = {  # supplier -> {sku: qty}; whole plan approved at engine quantities; SKU006, SKU007 paused
    "NutriFlow Distributors": {"SKU001": 750, "SKU013": 120},
    "GrainPath Logistics": {"SKU003": 300, "SKU004": 150},
    "Meridian Trade Co.": {"SKU011": 500},
}
NEVER = {"SKU006", "SKU007"}

d = Path(sys.argv[1])
pdf = next(d.glob("*.pdf"))
xlsx = next(d.glob("*.xlsx"))
csv_f = next(d.glob("*.csv"))
zf = next(d.glob("*.zip"))

reader = PdfReader(str(pdf))
assert len(reader.pages) == 3, len(reader.pages)
text = "\n".join(p.extract_text() for p in reader.pages)
for sup in EXPECTED:
    assert sup in text, sup
assert not any(s in text for s in NEVER)
assert "Approved by" in text and "Raghav Kulkarni" in text

wb = load_workbook(xlsx)
assert wb.sheetnames[0] == "Summary" and wb.sheetnames[-1] == "Approval Log"
for sup, lines in EXPECTED.items():
    ws = wb[sup[:31]]
    got = {r[0]: r[3] for r in ws.iter_rows(min_row=6, values_only=True) if r[0] and str(r[0]).startswith("SKU")}
    assert got == lines, (sup, got)

rows = list(csv.DictReader(io.StringIO(csv_f.read_text())))
got: dict = {}
for r in rows:
    got.setdefault(r["Vendor Name"], {})[r["SKU"]] = int(r["Quantity"])
assert got == EXPECTED, got

z = zipfile.ZipFile(zf)
assert z.testzip() is None
names = z.namelist()
assert "all_pos.pdf" in names and "all_pos.xlsx" in names and "zoho_import_all_pos.csv" in names
assert len([n for n in names if n.startswith("pdf/")]) == 3
print("PO FILES OK")
