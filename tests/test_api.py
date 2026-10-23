"""API, whole-plan approval, PO and ERPNext sync tests. Throwaway SQLite DB and a mock ERPNext; never Neon or a real ERP."""
import csv
import io
import json
import zipfile

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from planner_api import db, main
from planner_api.main import app
from planner_api.settings import get_settings
from reorder.io import read_dir

ANALYST = {"X-Role": "analyst"}
RAGHAV = {"X-Role": "raghav"}
PROC = {"X-Role": "procurement"}
EXPECTED = {"DIST001": {"SKU001": 750, "SKU013": 120}, "DIST002": {"SKU003": 300, "SKU004": 150},
            "DIST004": {"SKU011": 500}}
SUPPLIERS = ["NutriFlow Distributors", "GrainPath Logistics", "Meridian Trade Co."]


@pytest.fixture()
def client(tmp_path, monkeypatch):
    db.init_engine(f"sqlite:///{tmp_path / 'test.db'}")
    monkeypatch.setattr(get_settings(), "gemini_api_key", "")  # template explanations, no network
    for k in ("ERP_SYNC_ENABLED", "ERPNEXT_URL", "ERPNEXT_API_KEY", "ERPNEXT_API_SECRET", "ERPNEXT_COMPANY",
              "ERPNEXT_WAREHOUSE", "ERPNEXT_SUBMIT"):
        monkeypatch.delenv(k, raising=False)
    return TestClient(app)


class FakeERPNext:
    """Minimal ERPNext REST API: Supplier, Item and Purchase Order resources."""

    def __init__(self, missing_supplier=None):
        self.created, self.missing_supplier = [], missing_supplier

    def __call__(self, req: httpx.Request) -> httpx.Response:
        assert req.headers["Authorization"] == "token key1:secret1"
        path = req.url.path
        if req.method == "GET":
            filters = json.loads(req.url.params["filters"])
            if path.endswith("/Supplier"):
                name = filters[0][2]
                ok = name in SUPPLIERS and name != self.missing_supplier
                return httpx.Response(200, json={"data": [{"name": name, "supplier_name": name, "disabled": 0}] if ok else []})
            if path.endswith("/Item"):
                code = filters[0][2]
                return httpx.Response(200, json={"data": [{"name": code, "item_code": code, "disabled": 0}]})
            if path.endswith("/Purchase Order"):
                title = filters[0][2]
                hits = [{"name": c["_name"], "title": c["title"], "status": "Draft", "docstatus": 0}
                        for c in self.created if c["title"] == title]
                return httpx.Response(200, json={"data": hits})
        if req.method == "POST" and path.endswith("/Purchase Order"):
            doc = json.loads(req.content)
            doc["_name"] = f"PUR-ORD-2026-{len(self.created) + 1:05d}"
            self.created.append(doc)
            return httpx.Response(200, json={"data": {"name": doc["_name"], "status": "Draft", "docstatus": 0}})
        return httpx.Response(404, json={"exc_type": "DoesNotExistError"})


def connect_erp(monkeypatch, fake):
    for k, v in {"ERP_SYNC_ENABLED": "true", "ERPNEXT_URL": "https://freshroots.example.erpnext.com",
                 "ERPNEXT_API_KEY": "key1", "ERPNEXT_API_SECRET": "secret1", "ERPNEXT_COMPANY": "Freshroots"}.items():
        monkeypatch.setenv(k, v)
    monkeypatch.setattr(main, "_erp_transport", httpx.MockTransport(fake))


def upload(client, oct_dir, run_date="2026-10-19", headers=ANALYST, drop=None):
    files = [("files", (n, b, "text/csv")) for n, b in read_dir(oct_dir).items() if n != drop]
    return client.post("/api/runs", files=files, data={"run_date": run_date}, headers=headers)


def test_roles(client, oct_dir):
    assert upload(client, oct_dir, headers=RAGHAV).status_code == 403   # Raghav doesn't run plans
    assert upload(client, oct_dir, headers=PROC).status_code == 403
    assert client.get("/api/runs").status_code == 401
    run_id = upload(client, oct_dir).json()["run_id"]
    assert client.post(f"/api/runs/{run_id}/approve", headers=ANALYST).status_code == 403
    assert client.post(f"/api/runs/{run_id}/approve", headers=PROC).status_code == 403


def test_fatal_returns_exact_messages(client, oct_dir):
    r = upload(client, oct_dir, run_date="2026-10-30")
    assert r.status_code == 422
    assert "I7: inventory snapshot is stale or future-dated: 2026-10-19" in r.json()["validation"]["errors"]
    r = upload(client, oct_dir, drop="calendar.csv")
    assert "I1: missing file or column: calendar.csv" in r.json()["validation"]["errors"]


def test_fetch_latest_sap_data(client, oct_dir):
    r = client.post("/api/runs", data={"source": "sap_latest", "run_date": "2026-10-19"}, headers=ANALYST)
    assert r.status_code == 409  # nothing uploaded yet
    first = upload(client, oct_dir).json()["run_id"]
    info = client.get("/api/sap/latest", headers=ANALYST).json()
    assert info["run_id"] == first and len(info["files"]) == 5
    r = client.post("/api/runs", data={"source": "sap_latest", "run_date": "2026-10-19"}, headers=ANALYST)
    assert r.status_code == 200, r.text
    d = client.get(f"/api/runs/{r.json()['run_id']}", headers=ANALYST).json()
    assert d["source"] == "sap_latest" and d["source_run_id"] == first
    assert d["summary"]["total_units_now"] == 1820
    assert [f["sha256"] for f in d["files"]] == [f["sha256"] for f in info["files"]]


def test_reject_needs_note_and_makes_no_pos(client, oct_dir):
    run_id = upload(client, oct_dir).json()["run_id"]
    assert client.post(f"/api/runs/{run_id}/reject", headers=RAGHAV, json={"note": " "}).status_code == 422
    r = client.post(f"/api/runs/{run_id}/reject", headers=RAGHAV, json={"note": "Festival uplift looks double-counted"})
    assert r.json()["status"] == "rejected"
    assert client.get(f"/api/runs/{run_id}/pos", headers=PROC).json() == []
    assert client.post(f"/api/runs/{run_id}/approve", headers=RAGHAV).status_code == 409


def test_approve_without_erp_keeps_pos_and_reason(client, oct_dir):
    run_id = upload(client, oct_dir).json()["run_id"]
    pos = client.post(f"/api/runs/{run_id}/approve", headers=RAGHAV).json()["pos"]
    assert {p["distributor_id"]: {l["sku_id"]: l["quantity"] for l in p["lines"]} for p in pos} == EXPECTED
    assert all(p["status"] == "Generated" and "ERPNext not connected" in p["erp_error"] for p in pos)


def test_approve_creates_pos_in_erpnext(client, oct_dir, monkeypatch):
    fake = FakeERPNext()
    connect_erp(monkeypatch, fake)
    run_id = upload(client, oct_dir).json()["run_id"]
    r = client.post(f"/api/runs/{run_id}/approve", headers=RAGHAV, json={"note": "Looks right"})
    assert r.status_code == 200, r.text
    pos = r.json()["pos"]
    assert all(p["status"] == "Entered in ERPNext" and p["erp_po_name"].startswith("PUR-ORD-") and p["erp_error"] is None
               for p in pos)
    assert all(p["erp_url"] == f"https://freshroots.example.erpnext.com/app/purchase-order/{p['erp_po_name']}" for p in pos)
    # ERPNext received exactly the approved quantities, one PO per supplier, no rates, our PO number as title
    got = {c["supplier"]: {l["item_code"]: l["qty"] for l in c["items"]} for c in fake.created}
    assert got == {"NutriFlow Distributors": {"SKU001": 750, "SKU013": 120},
                   "GrainPath Logistics": {"SKU003": 300, "SKU004": 150},
                   "Meridian Trade Co.": {"SKU011": 500}}
    assert all("rate" not in l for c in fake.created for l in c["items"])
    assert {c["title"] for c in fake.created} == {p["po_number"] for p in pos}
    assert all(c["company"] == "Freshroots" and c["schedule_date"] == c["transaction_date"] and "docstatus" not in c
               for c in fake.created)
    audit = client.get(f"/api/audit?run_id={run_id}", headers=PROC).json()
    assert sum(a["action"] == "po.erp_created" for a in audit) == 3
    # manual "entered" is blocked once ERPNext has it
    assert client.patch(f"/api/pos/{pos[0]['id']}", headers=PROC,
                        json={"status": "Entered in ERPNext", "erp_ref": "x"}).status_code == 409


def test_erp_failure_then_retry(client, oct_dir, monkeypatch):
    fake = FakeERPNext(missing_supplier="GrainPath Logistics")
    connect_erp(monkeypatch, fake)
    run_id = upload(client, oct_dir).json()["run_id"]
    pos = {p["distributor_id"]: p for p in client.post(f"/api/runs/{run_id}/approve", headers=RAGHAV).json()["pos"]}
    assert pos["DIST001"]["status"] == "Entered in ERPNext"
    assert pos["DIST002"]["status"] == "ERPNext sync failed"
    assert "Supplier 'GrainPath Logistics' not found" in pos["DIST002"]["erp_error"]
    assert client.post(f"/api/pos/{pos['DIST002']['id']}/erp-sync", headers=RAGHAV).status_code == 403
    fake.missing_supplier = None  # supplier created in ERPNext
    again = client.post(f"/api/pos/{pos['DIST002']['id']}/erp-sync", headers=PROC).json()
    assert again["status"] == "Entered in ERPNext"
    # a retry of an already-synced PO doesn't create a duplicate
    client.post(f"/api/pos/{pos['DIST001']['id']}/erp-sync", headers=PROC)
    assert len(fake.created) == 3


def test_manual_entry_when_erp_unavailable(client, oct_dir):
    run_id = upload(client, oct_dir).json()["run_id"]
    po = client.post(f"/api/runs/{run_id}/approve", headers=RAGHAV).json()["pos"][0]
    assert client.patch(f"/api/pos/{po['id']}", headers=ANALYST, json={"status": "Entered in ERPNext", "erp_ref": "PO-9"}).status_code == 403
    p = client.patch(f"/api/pos/{po['id']}", headers=PROC, json={"status": "Entered in ERPNext", "erp_ref": "PO-9"}).json()
    assert (p["status"], p["erp_ref"], p["erp_po_name"]) == ("Entered in ERPNext", "PO-9", None)


def test_po_files(client, oct_dir):
    run_id = upload(client, oct_dir).json()["run_id"]
    pos = client.post(f"/api/runs/{run_id}/approve", headers=RAGHAV).json()["pos"]
    for p in pos:
        assert p["total_qty"] == sum(l["quantity"] for l in p["lines"])
        assert p["po_number"].startswith("FR-PO-") and p["po_number"].endswith(f"-{p['distributor_id']}-001")
    assert not {l["sku_id"] for p in pos for l in p["lines"]} & {"SKU006", "SKU007"}
    pdf = client.get(f"/api/runs/{run_id}/pos/download?format=pdf", headers=PROC)
    reader = PdfReader(io.BytesIO(pdf.content))
    assert len(reader.pages) == 3
    text = "".join(pg.extract_text() for pg in reader.pages)
    assert "Raghav Kulkarni" in text and "Generated by Freshroots Reorder Planner. Created in ERPNext on approval." in text
    wb = load_workbook(io.BytesIO(client.get(f"/api/runs/{run_id}/pos/download?format=xlsx", headers=PROC).content))
    assert wb.sheetnames[0] == "Summary" and wb.sheetnames[-1] == "Approval Log" and len(wb.sheetnames) == 5
    rows = list(csv.DictReader(io.StringIO(client.get(f"/api/runs/{run_id}/pos/download?format=csv", headers=PROC).text)))
    assert sum(int(r["Quantity"]) for r in rows) == 1820
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/runs/{run_id}/pos/download?format=zip", headers=PROC).content))
    assert z.testzip() is None and "plan/reorder_plan_2026-10.csv" in z.namelist()


def test_plan_files_match_engine(client, oct_dir):
    run_id = upload(client, oct_dir).json()["run_id"]
    csv_text = client.get(f"/api/runs/{run_id}/plan?format=csv", headers=RAGHAV).text
    assert ",750,0,6.14,0,R3," in csv_text


def test_master_data_tables(client, oct_dir):
    assert client.get("/api/master", headers=PROC).json() is None
    run_id = upload(client, oct_dir).json()["run_id"]
    m = client.get("/api/master", headers=PROC).json()
    assert m["run_id"] == run_id
    assert len(m["skus"]) == 15 and len(m["suppliers"]) == 4 and len(m["calendar"]) == 24
    sku1 = next(k for k in m["skus"] if k["sku_id"] == "SKU001")
    assert (sku1["supplier_name"], sku1["units_on_hand"], sku1["standard_reorder_qty"]) == ("NutriFlow Distributors", 900, 500)
    assert [h["week"] for h in sku1["sales"]][-1] == "2026-W42" and len(sku1["sales"]) == 16
    assert {c["year_month"] for c in m["calendar"] if c["is_festival_month"]} >= {"2026-03", "2026-10", "2026-11"}
    # a run on the latest SAP data doesn't duplicate master rows
    client.post("/api/runs", data={"source": "sap_latest", "run_date": "2026-10-19"}, headers=ANALYST)
    with db._Session() as s:
        from sqlalchemy import func, select
        assert s.scalar(select(func.count()).select_from(db.SkuMaster)) == 15
        assert s.scalar(select(func.count()).select_from(db.WeeklySales)) == 240


def test_backfill_master_data(client, oct_dir):
    run_id = upload(client, oct_dir).json()["run_id"]
    with db._Session() as s:
        from sqlalchemy import delete
        for t in (db.SkuMaster, db.DistributorMaster, db.WeeklySales, db.CurrentInventory, db.FestivalCalendar):
            s.execute(delete(t))
        s.commit()
    db._backfill_master_data()
    assert client.get("/api/master", headers=PROC).json()["run_id"] == run_id
    assert len(client.get("/api/master", headers=PROC).json()["skus"]) == 15
